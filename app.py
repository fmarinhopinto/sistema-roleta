import io
import os
import re
import unicodedata
from pathlib import Path
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(
    page_title="Sistema Roleta - Controle Comercial",
    page_icon="🎯",
    layout="wide"
)

# --- DIRETÓRIOS LOCAIS ---
BASE_DIR = Path(__file__).resolve().parent
PASTA_ENTRADA = BASE_DIR / "arquivos_entrada"
PASTA_SAIDA = BASE_DIR / "planilhas_geradas"

PASTA_ENTRADA.mkdir(parents=True, exist_ok=True)
PASTA_SAIDA.mkdir(parents=True, exist_ok=True)

st.title("🎯 Sistema Roleta - Controle Comercial")
st.caption(f"Diretório base: `{BASE_DIR}`")

# --- FUNÇÕES DE SUPORTE ---
def normalizar_texto(texto):
    if pd.isna(texto):
        return ""
    texto = str(texto).strip().upper()
    texto = unicodedata.normalize('NFKD', texto).encode('ASCII', 'ignore').decode('utf-8')
    texto = re.sub(r'[\.\-\/\,\_]', ' ', texto)
    return " ".join(texto.split())

@st.cache_data(show_spinner=False)
def carregar_dados(origem):
    nome = origem.name if hasattr(origem, 'name') else str(origem)
    if nome.lower().endswith('.csv'):
        try:
            return pd.read_csv(origem, sep=';', encoding='latin1')
        except Exception:
            if hasattr(origem, 'seek'):
                origem.seek(0)
            return pd.read_csv(origem)
    else:
        return pd.read_excel(origem)

def identificar_coluna(colunas, possiveis_nomes):
    for nome in possiveis_nomes:
        nome_norm = normalizar_texto(nome).lower()
        for col in colunas:
            col_norm = normalizar_texto(col).lower()
            if (nome_norm in col_norm) or (nome.lower() in str(col).lower()):
                return col
    return colunas[0] if len(colunas) > 0 else None

def buscar_arquivo_padrao(lista_arquivos, termos):
    for f in lista_arquivos:
        f_norm = normalizar_texto(f).lower()
        if any(t in f_norm for t in termos):
            return f
    return "Selecione..."

def limpar_valor(val):
    if pd.isna(val):
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace('R$', '').strip()
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        s = s.replace(',', '.')
    return pd.to_numeric(s, errors='coerce') or 0.0

def estilizar_aba(ws, df_dados, col_id_nome, col_id_val, col_cli_nome, col_data_dt, titulo_col_id="PEDIDO", titulo_status=None):
    """
    Formata e preenche uma aba no padrão exato da Roleta:
    - Colunas A-G: Analítico
    - Colunas I-K: Resumo por Vendedor com fórmulas =SUM()
    """
    fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="000000")
    fill_cabecalho = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    borda_fina = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fonte_alerta_vermelho = Font(name="Calibri", size=11, bold=True, color="9C0006")
    fill_alerta_vermelho = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # Cabeçalho Principal (A1:G1)
    cabecalhos_dados = ['VENDEDOR', titulo_col_id, 'CLIENTE', 'VALOR', 'STATUS', 'DATA', 'BRINDE ROLETA']
    for idx, h in enumerate(cabecalhos_dados, 1):
        cell = ws.cell(row=1, column=idx, value=h)
        cell.font = fonte_cabecalho
        cell.fill = fill_cabecalho
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Cabeçalho Lateral (I1:K1)
    label_qtd = 'NUMERO DE PEDIDOS' if titulo_col_id == "PEDIDO" else 'NUMERO DE PROPOSTAS'
    cabecalhos_resumo = ['VENDEDOR', label_qtd, 'TOTAL']
    for idx, h in enumerate(cabecalhos_resumo, 9):
        cell = ws.cell(row=1, column=idx, value=h)
        cell.font = fonte_cabecalho
        cell.fill = fill_cabecalho
        cell.alignment = Alignment(horizontal="center", vertical="center")

    grupos_vendedores = {}
    linha_atual = 2

    for _, r in df_dados.iterrows():
        vend = r['VENDEDOR_CLEAN']
        tem_duas_compras = bool(r.get('TEM_DUAS_COMPRAS', False))
        obs_brinde = str(r.get('BRINDE ROLETA', '')) if pd.notna(r.get('BRINDE ROLETA', '')) else ""

        ws.cell(row=linha_atual, column=1, value=vend)
        ws.cell(row=linha_atual, column=2, value=r[col_id_nome])
        ws.cell(row=linha_atual, column=3, value=str(r[col_cli_nome]))

        c_val = ws.cell(row=linha_atual, column=4, value=r['VALOR_NUM'])
        c_val.number_format = 'R$ #,##0.00'

        status_val = r['STATUS'] if 'STATUS' in r else (titulo_status or "")
        ws.cell(row=linha_atual, column=5, value=status_val)

        c_dt = ws.cell(row=linha_atual, column=6, value=r[col_data_dt].date() if pd.notna(r[col_data_dt]) else "")
        c_dt.number_format = 'DD/MM/YYYY'
        c_dt.alignment = Alignment(horizontal="center")

        ws.cell(row=linha_atual, column=7, value=obs_brinde)

        for c in range(1, 8):
            cell_item = ws.cell(row=linha_atual, column=c)
            cell_item.border = borda_fina
            if tem_duas_compras:
                cell_item.fill = fill_alerta_vermelho
                cell_item.font = fonte_alerta_vermelho
                if c == 7:
                    cell_item.alignment = Alignment(horizontal="center", vertical="center")

        if vend not in grupos_vendedores:
            grupos_vendedores[vend] = []
        grupos_vendedores[vend].append(linha_atual)

        linha_atual += 1

    # Tabela Resumo (Colunas I, J, K)
    linha_resumo = 2
    for vend, linhas in sorted(grupos_vendedores.items()):
        primeira = linhas[0]
        ultima = linhas[-1]
        qtd = len(linhas)

        ws.cell(row=linha_resumo, column=9, value=vend)
        ws.cell(row=linha_resumo, column=10, value=qtd).alignment = Alignment(horizontal="center")

        c_tot = ws.cell(row=linha_resumo, column=11, value=f"=SUM(D{primeira}:D{ultima})")
        c_tot.number_format = 'R$ #,##0.00'

        for c in range(9, 12):
            ws.cell(row=linha_resumo, column=c).border = borda_fina

        linha_resumo += 1

    # Auto-ajuste de colunas
    for col in ws.columns:
        letra = get_column_letter(col[0].column)
        tamanho_max = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[letra].width = max(tamanho_max + 3, 12)

def estilizar_aba_propostas(ws, df_resumo_adm):
    """
    Formata e preenche a Aba 2 no padrão exato da Roleta (Tabela Dinâmica ADM):
    - Coluna A: Vendedor
    - Coluna B: Qtd
    Cabeçalho com preenchimento amarelo (#FFFF00), negrito e bordas finas.
    """
    fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="000000")
    fill_cabecalho = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    borda_fina = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fonte_dados = Font(name="Calibri", size=11, bold=False, color="000000")

    # Cabeçalho (A1:B1)
    c1 = ws.cell(row=1, column=1, value="Vendedor")
    c1.font = fonte_cabecalho
    c1.fill = fill_cabecalho
    c1.border = borda_fina
    c1.alignment = Alignment(horizontal="left", vertical="center")

    c2 = ws.cell(row=1, column=2, value="Qtd")
    c2.font = fonte_cabecalho
    c2.fill = fill_cabecalho
    c2.border = borda_fina
    c2.alignment = Alignment(horizontal="center", vertical="center")

    # Linhas de dados
    for idx, r in enumerate(df_resumo_adm.itertuples(), start=2):
        c_vend = ws.cell(row=idx, column=1, value=str(r.Vendedor))
        c_vend.font = fonte_dados
        c_vend.border = borda_fina
        c_vend.alignment = Alignment(horizontal="left", vertical="center")

        c_qtd = ws.cell(row=idx, column=2, value=int(r.Qtd))
        c_qtd.font = fonte_dados
        c_qtd.border = borda_fina
        c_qtd.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-ajuste de largura das colunas
    max_len_v = max([len(str(r.Vendedor)) for r in df_resumo_adm.itertuples()] + [8])
    ws.column_dimensions['A'].width = max(max_len_v + 4, 22)
    ws.column_dimensions['B'].width = 12

def estilizar_aba_produtos(ws, df_top_notas, df_lideres_cat):
    """
    Formata e preenche a Aba 3 de Destaques da Planilha de Produtos:
    - Bloco 1: Top 10 Maiores Vendas em Nota Única
    - Bloco 2: Líderes de Vendas por Categoria
    """
    fonte_titulo = Font(name="Calibri", size=13, bold=True, color="1F4E79")
    fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="000000")
    fill_cab_azul = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_cab_amarelo = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    borda_fina = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fonte_dados = Font(name="Calibri", size=11, bold=False, color="000000")
    fonte_destaque = Font(name="Calibri", size=11, bold=True, color="000000")

    # --- BLOCO 1: TOP NOTAS ÚNICAS ---
    ws.cell(row=1, column=1, value="🏆 TOP MAIORES VENDAS EM NOTA ÚNICA").font = fonte_titulo

    headers_1 = ["POSIÇÃO", "VENDEDOR", "NOTA", "VALOR TOTAL"]
    for idx, h in enumerate(headers_1, 1):
        c = ws.cell(row=2, column=idx, value=h)
        c.font = fonte_cabecalho
        c.fill = fill_cab_azul
        c.border = borda_fina
        c.alignment = Alignment(horizontal="center", vertical="center")

    linha_atual = 3
    for pos, r in enumerate(df_top_notas.head(10).itertuples(), start=1):
        c_pos = ws.cell(row=linha_atual, column=1, value=f"{pos}º")
        c_pos.alignment = Alignment(horizontal="center")
        c_pos.font = fonte_destaque if pos == 1 else fonte_dados
        c_pos.border = borda_fina

        c_vend = ws.cell(row=linha_atual, column=2, value=str(r.Vendedor))
        c_vend.font = fonte_destaque if pos == 1 else fonte_dados
        c_vend.border = borda_fina

        nota_val = int(r.Nota) if pd.notna(r.Nota) and isinstance(r.Nota, (int, float)) and r.Nota == int(r.Nota) else str(r.Nota)
        c_nota = ws.cell(row=linha_atual, column=3, value=nota_val)
        c_nota.alignment = Alignment(horizontal="center")
        c_nota.font = fonte_dados
        c_nota.border = borda_fina

        c_val = ws.cell(row=linha_atual, column=4, value=float(r.Valor_Total))
        c_val.number_format = 'R$ #,##0.00'
        c_val.font = fonte_destaque if pos == 1 else fonte_dados
        c_val.border = borda_fina

        linha_atual += 1

    # Espaço entre blocos
    linha_atual += 2

    # --- BLOCO 2: LÍDERES POR CATEGORIA ---
    ws.cell(row=linha_atual, column=1, value="🏷️ MAIOR VENDA POR CATEGORIA DE PRODUTO").font = fonte_titulo
    linha_atual += 1

    headers_2 = ["CATEGORIA", "VENDEDOR LÍDER", "TOTAL VENDIDO"]
    for idx, h in enumerate(headers_2, 1):
        c = ws.cell(row=linha_atual, column=idx, value=h)
        c.font = fonte_cabecalho
        c.fill = fill_cab_amarelo
        c.border = borda_fina
        c.alignment = Alignment(horizontal="center", vertical="center")

    linha_atual += 1
    for r in df_lideres_cat.itertuples():
        c_cat = ws.cell(row=linha_atual, column=1, value=str(r.Categoria))
        c_cat.font = fonte_dados
        c_cat.border = borda_fina

        c_vend = ws.cell(row=linha_atual, column=2, value=str(r.Vendedor))
        c_vend.font = fonte_destaque
        c_vend.border = borda_fina

        c_val = ws.cell(row=linha_atual, column=3, value=float(r.Valor_Total))
        c_val.number_format = 'R$ #,##0.00'
        c_val.font = fonte_dados
        c_val.border = borda_fina

        linha_atual += 1

    # Auto-ajuste de largura de colunas
    for col in ws.columns:
        letra = get_column_letter(col[0].column)
        tamanho_max = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[letra].width = max(tamanho_max + 3, 14)


CATEGORIAS_ROLETA = [
    "ALIMENTOS E BEBIDAS",
    "EMBALAGENS E DESCARTAVEIS",
    "EPI'S",
    "HIGIENE E LIMPEZA",
    "INFORMÁTICA",
    "MRO",
    "OFFICE"
]

def extrair_ranking_categoria(df_cat_raw, categoria_alvo, df_fortpel=None, df_adm=None):
    """
    Processa a planilha de categorias (formato wide como categorias.xlsx ou long como produtos.xlsx)
    e retorna: (df_ranking_cat, total_geral_cat, df_lideres_todas_categorias)
    """
    if df_cat_raw is None or df_cat_raw.empty:
        return pd.DataFrame(), 0.0, pd.DataFrame()

    # Monta mapa de cliente -> vendedor a partir do Fortpel e ADM
    mapa_cli_vend = {}
    if df_fortpel is not None and not df_fortpel.empty:
        col_c = None
        for cand in ['nome_cliente', 'cliente', 'razao', 'nome']:
            for c in df_fortpel.columns:
                if cand in str(c).lower():
                    col_c = c
                    break
            if col_c:
                break

        col_v = None
        for cand in ['carteira', 'operador', 'vendedor', 'gestor']:
            for c in df_fortpel.columns:
                if cand in str(c).lower():
                    col_v = c
                    break
            if col_v:
                break

        if col_c and col_v:
            df_f_clean = df_fortpel.dropna(subset=[col_c, col_v]).copy()
            df_f_clean['__C_NORM'] = df_f_clean[col_c].apply(normalizar_texto)
            grp = df_f_clean.groupby('__C_NORM')[col_v].agg(lambda x: x.mode()[0] if not x.mode().empty else x.iloc[0])
            for cli, v in grp.items():
                mapa_cli_vend[cli] = str(v).strip().upper()

    if df_adm is not None and not df_adm.empty:
        col_ac = None
        col_av = None
        for cand in ['nome do cliente', 'cliente', 'nome']:
            for c in df_adm.columns:
                if cand in str(c).lower():
                    col_ac = c
                    break
            if col_ac:
                break
        for cand in ['vendedor do cliente', 'vendedor', 'operador']:
            for c in df_adm.columns:
                if cand in str(c).lower():
                    col_av = c
                    break
            if col_av:
                break
        if col_ac and col_av:
            for _, r in df_adm.dropna(subset=[col_ac, col_av]).iterrows():
                cn = normalizar_texto(r[col_ac])
                if cn and cn not in mapa_cli_vend:
                    mapa_cli_vend[cn] = str(r[col_av]).strip().upper()

    # Trata cabeçalhos da planilha de categorias
    df_work = df_cat_raw.copy()
    if str(df_work.iloc[0, 0]).strip().lower() in ['cliente', 'codigo', 'cod']:
        df_work = df_work.iloc[1:].copy()

    # Remove linhas finais de total / filtros
    df_work = df_work[~df_work.iloc[:, 0].astype(str).str.upper().str.contains('TOTAL|FILTROS APLICADOS', na=False)].copy()

    # Verifica se é formato WIDE (onde as colunas são as próprias categorias)
    cols_map = {}
    for col in df_work.columns:
        c_norm = normalizar_texto(col).replace("'", "").replace(" ", "")
        for cat in CATEGORIAS_ROLETA:
            cat_norm = normalizar_texto(cat).replace("'", "").replace(" ", "")
            if cat_norm == c_norm or cat_norm in c_norm or c_norm in cat_norm:
                cols_map[cat] = col

    lideres_todas = []

    if len(cols_map) >= 2:
        # FORMATO WIDE
        col_id_primeira = df_work.columns[0]
        df_work['__ID_NORM'] = df_work[col_id_primeira].apply(normalizar_texto)

        def obter_vendedor(id_norm):
            if id_norm in mapa_cli_vend:
                return mapa_cli_vend[id_norm]
            return id_norm

        df_work['__VENDEDOR'] = df_work['__ID_NORM'].apply(obter_vendedor)

        for cat in CATEGORIAS_ROLETA:
            c_nome = cols_map.get(cat)
            if c_nome and c_nome in df_work.columns:
                df_work[f'__VAL_{cat}'] = df_work[c_nome].apply(limpar_valor)
                res_cat = (
                    df_work.groupby('__VENDEDOR', as_index=False)[f'__VAL_{cat}']
                    .sum()
                    .rename(columns={'__VENDEDOR': 'Vendedor', f'__VAL_{cat}': 'Valor_Total'})
                    .sort_values(by='Valor_Total', ascending=False)
                )
                res_cat = res_cat[res_cat['Valor_Total'] > 0]
                if not res_cat.empty:
                    top_v = res_cat.iloc[0]
                    lideres_todas.append({
                        'Categoria': cat,
                        'Vendedor': top_v['Vendedor'],
                        'Valor_Total': top_v['Valor_Total'],
                        'Total_Categoria': res_cat['Valor_Total'].sum()
                    })

        col_alvo = cols_map.get(categoria_alvo)
        if col_alvo and col_alvo in df_work.columns:
            df_work['__VAL_ALVO'] = df_work[col_alvo].apply(limpar_valor)
            df_rank_alvo = (
                df_work.groupby('__VENDEDOR', as_index=False)
                .agg(
                    Total_Vendido=('__VAL_ALVO', 'sum'),
                    Qtd_Itens=('__VAL_ALVO', lambda s: (s > 0).sum())
                )
                .rename(columns={'__VENDEDOR': 'Vendedor'})
                .sort_values(by='Total_Vendido', ascending=False)
            )
            df_rank_alvo = df_rank_alvo[df_rank_alvo['Total_Vendido'] > 0]
            total_alvo = df_rank_alvo['Total_Vendido'].sum()
            return df_rank_alvo, total_alvo, pd.DataFrame(lideres_todas)

    else:
        # FORMATO LONG (colunas Categoria, Vendedor, Valor)
        col_c = identificar_coluna(df_work.columns, ["categoria", "grupo", "linha", "subgrupo", "descr. produto", "descricao"])
        col_v = identificar_coluna(df_work.columns, ["vendedor", "operador", "gestor", "representante", "carteira"])
        col_val = identificar_coluna(df_work.columns, ["vendas total", "venda total", "venda", "valor total", "valor", "total", "preco"])

        if col_c and col_v and col_val:
            df_work['__VALOR_NUM'] = df_work[col_val].apply(limpar_valor)
            df_work['__VEND_LIMPO'] = df_work[col_v].astype(str).str.strip().str.upper()

            for cat in CATEGORIAS_ROLETA:
                cat_norm = normalizar_texto(cat)
                m_cat = df_work[col_c].apply(lambda x: cat_norm in normalizar_texto(x) or normalizar_texto(x) in cat_norm)
                df_sub = df_work[m_cat]
                if not df_sub.empty:
                    grp = (
                        df_sub.groupby('__VEND_LIMPO', as_index=False)['__VALOR_NUM']
                        .sum()
                        .rename(columns={'__VEND_LIMPO': 'Vendedor', '__VALOR_NUM': 'Valor_Total'})
                        .sort_values(by='Valor_Total', ascending=False)
                    )
                    grp = grp[grp['Valor_Total'] > 0]
                    if not grp.empty:
                        lideres_todas.append({
                            'Categoria': cat,
                            'Vendedor': grp.iloc[0]['Vendedor'],
                            'Valor_Total': grp.iloc[0]['Valor_Total'],
                            'Total_Categoria': grp['Valor_Total'].sum()
                        })

            cat_alvo_norm = normalizar_texto(categoria_alvo)
            m_alvo = df_work[col_c].apply(lambda x: cat_alvo_norm in normalizar_texto(x) or normalizar_texto(x) in cat_alvo_norm)
            df_sub_alvo = df_work[m_alvo]
            if not df_sub_alvo.empty:
                df_rank_alvo = (
                    df_sub_alvo.groupby('__VEND_LIMPO', as_index=False)
                    .agg(
                        Total_Vendido=('__VALOR_NUM', 'sum'),
                        Qtd_Itens=('__VALOR_NUM', 'count')
                    )
                    .rename(columns={'__VEND_LIMPO': 'Vendedor'})
                    .sort_values(by='Total_Vendido', ascending=False)
                )
                df_rank_alvo = df_rank_alvo[df_rank_alvo['Total_Vendido'] > 0]
                total_alvo = df_rank_alvo['Total_Vendido'].sum()
                return df_rank_alvo, total_alvo, pd.DataFrame(lideres_todas)

    return pd.DataFrame(), 0.0, pd.DataFrame(lideres_todas)

def extrair_maior_venda_nota(df_fortpel, col_nota, col_vend, col_val, dt_inicio=None, dt_fim=None, df_cat_raw=None):
    """
    Calcula as maiores vendas em nota única.
    Extrai primariamente a partir da base de pedidos do Fortpel.
    Se df_cat_raw contiver itens detalhados por nota (ex: produtos.xlsx com n_nota_obs), complementa/reforça.
    """
    df_top_notas = pd.DataFrame()

    # 1. Tenta extrair de df_fortpel
    if df_fortpel is not None and not df_fortpel.empty and col_nota and col_vend and col_val:
        df_f = df_fortpel.copy()
        df_f['__VALOR_NUM'] = df_f[col_val].apply(limpar_valor)
        df_f['__VEND_LIMPO'] = df_f[col_vend].astype(str).str.strip().str.upper()

        if '__DATA_DT' in df_f.columns and pd.notna(dt_inicio) and pd.notna(dt_fim):
            mask_dt = (df_f['__DATA_DT'].dt.date >= dt_inicio) & (df_f['__DATA_DT'].dt.date <= dt_fim)
            df_sub = df_f[mask_dt].copy()
            if df_sub.empty:
                df_sub = df_f.copy()
        else:
            df_sub = df_f.copy()

        # Group by nota e vendedor
        df_grp = (
            df_sub.groupby([col_nota, '__VEND_LIMPO'], as_index=False)['__VALOR_NUM']
            .sum()
            .rename(columns={col_nota: 'Nota', '__VEND_LIMPO': 'Vendedor', '__VALOR_NUM': 'Valor_Total'})
            .sort_values(by='Valor_Total', ascending=False)
        )
        df_grp = df_grp[df_grp['Valor_Total'] > 0]
        if df_grp['Valor_Total'].nunique() > 1:
            df_top_notas = df_grp

    # 2. Se df_fortpel tiver valores constantes ou vazios, verifica se df_cat_raw tem itens por nota (como produtos.xlsx)
    if df_top_notas.empty and df_cat_raw is not None and not df_cat_raw.empty:
        col_n = identificar_coluna(df_cat_raw.columns, ["n_nota_obs", "nota", "nf", "pedido"])
        col_v = identificar_coluna(df_cat_raw.columns, ["vendedor", "operador", "gestor", "carteira"])
        col_vl = identificar_coluna(df_cat_raw.columns, ["vendas total", "venda total", "venda", "valor total", "valor"])
        if col_n and col_v and col_vl and col_n in df_cat_raw.columns and col_v in df_cat_raw.columns and col_vl in df_cat_raw.columns:
            df_c_clean = df_cat_raw.dropna(subset=[col_n, col_v, col_vl]).copy()
            df_c_clean['__VAL_N'] = df_c_clean[col_vl].apply(limpar_valor)
            df_c_clean['__VEND_N'] = df_c_clean[col_v].astype(str).str.strip().str.upper()
            df_top_notas = (
                df_c_clean.groupby([col_n, '__VEND_N'], as_index=False)['__VAL_N']
                .sum()
                .rename(columns={col_n: 'Nota', '__VEND_N': 'Vendedor', '__VAL_N': 'Valor_Total'})
                .sort_values(by='Valor_Total', ascending=False)
            )

    return df_top_notas


# --- SELEÇÃO / UPLOAD DOS ARQUIVOS ---
st.subheader("1. Arquivos de Entrada (Bases BI, Fortpel, ADM e Categorias)")

modo_leitura = st.radio(
    "Como deseja carregar os arquivos?",
    ["Fazer Upload pelo Navegador", f"Ler direto da pasta local ({PASTA_ENTRADA.name} ou raiz)"],
    horizontal=True
)

file_novos = None
file_reativados = None
file_fortpel = None
file_adm = None
file_categorias = None
categoria_solicitada = "HIGIENE E LIMPEZA"

if modo_leitura == "Fazer Upload pelo Navegador":
    c1, c2 = st.columns(2)
    c3, c4 = st.columns(2)
    with c1:
        file_novos = st.file_uploader("📥 1. BI Clientes Novos", type=["xlsx", "xls", "csv"], key="up_novos")
    with c2:
        file_reativados = st.file_uploader("📥 2. BI Clientes Reativados", type=["xlsx", "xls", "csv"], key="up_reat")
    with c3:
        file_fortpel = st.file_uploader("📥 3. Acompanhamento Fortpel (Pedidos)", type=["xlsx", "xls", "csv"], key="up_fort")
    with c4:
        file_adm = st.file_uploader("📥 4. Planilha ADM (Orçamentos Efetivados)", type=["xlsx", "xls", "csv"], key="up_adm")

    c5, c6 = st.columns([2, 1])
    with c5:
        file_categorias = st.file_uploader("📥 5. Planilha de Categorias do BI", type=["xlsx", "xls", "csv"], key="up_cat")
    with c6:
        categoria_solicitada = st.selectbox(
            "🏷️ Categoria Solicitada:",
            CATEGORIAS_ROLETA,
            index=3,
            help="Selecione a categoria para identificar o vendedor com a maior venda.",
            key="sel_cat_req"
        )

else:
    # Busca tanto na pasta 'arquivos_entrada' quanto na pasta raiz do sistema
    arquivos_locais = sorted(list(set(
        [f.name for f in PASTA_ENTRADA.iterdir() if f.is_file() and not f.name.startswith('~$') and f.suffix.lower() in ['.xlsx', '.xls', '.csv']] +
        [f.name for f in BASE_DIR.iterdir() if f.is_file() and not f.name.startswith('~$') and f.suffix.lower() in ['.xlsx', '.xls', '.csv']]
    )))

    def resolver_caminho(nome):
        if (PASTA_ENTRADA / nome).exists():
            return PASTA_ENTRADA / nome
        return BASE_DIR / nome

    if not arquivos_locais:
        st.info(f"Nenhum arquivo encontrado em `{PASTA_ENTRADA}` ou `{BASE_DIR}`. Coloque as planilhas lá ou use a opção de Upload.")
    else:
        padrao_novos = buscar_arquivo_padrao(arquivos_locais, ["novo"])
        padrao_reat = buscar_arquivo_padrao(arquivos_locais, ["reativ"])
        padrao_fort = buscar_arquivo_padrao(arquivos_locais, ["acomp", "fortpel"])
        padrao_adm = buscar_arquivo_padrao(arquivos_locais, ["adm", "orcamento", "proposta"])
        padrao_cat = buscar_arquivo_padrao(arquivos_locais, ["categoria", "cat", "prod", "data (7", "data (8"])

        c1, c2 = st.columns(2)
        c3, c4 = st.columns(2)
        opcoes = ["Selecione..."] + arquivos_locais

        with c1:
            idx_n = opcoes.index(padrao_novos) if padrao_novos in opcoes else 0
            nome_n = st.selectbox("1. BI Clientes Novos:", opcoes, index=idx_n, key="loc_novos")
            if nome_n != "Selecione...":
                file_novos = resolver_caminho(nome_n)

        with c2:
            idx_r = opcoes.index(padrao_reat) if padrao_reat in opcoes else 0
            nome_r = st.selectbox("2. BI Clientes Reativados:", opcoes, index=idx_r, key="loc_reat")
            if nome_r != "Selecione...":
                file_reativados = resolver_caminho(nome_r)

        with c3:
            idx_f = opcoes.index(padrao_fort) if padrao_fort in opcoes else 0
            nome_f = st.selectbox("3. Acompanhamento Fortpel:", opcoes, index=idx_f, key="loc_fort")
            if nome_f != "Selecione...":
                file_fortpel = resolver_caminho(nome_f)

        with c4:
            idx_a = opcoes.index(padrao_adm) if padrao_adm in opcoes else 0
            nome_a = st.selectbox("4. Planilha ADM (Orçamentos Efetivados):", opcoes, index=idx_a, key="loc_adm")
            if nome_a != "Selecione...":
                file_adm = resolver_caminho(nome_a)

        c5, c6 = st.columns([2, 1])
        with c5:
            idx_p = opcoes.index(padrao_cat) if padrao_cat in opcoes else 0
            nome_p = st.selectbox("5. Planilha de Categorias do BI (Opcional):", opcoes, index=idx_p, key="loc_cat")
            if nome_p != "Selecione...":
                file_categorias = resolver_caminho(nome_p)
        with c6:
            categoria_solicitada = st.selectbox(
                "🏷️ Categoria Solicitada:",
                CATEGORIAS_ROLETA,
                index=3,
                help="Selecione a categoria para identificar o vendedor com a maior venda.",
                key="sel_cat_req_loc"
            )


# --- MAPEAMENTO E PROCESSAMENTO ---
if file_novos and file_reativados and file_fortpel and file_adm:
    try:
        with st.spinner("Lendo arquivos Excel/CSV..."):
            df_novos_raw = carregar_dados(file_novos)
            df_reativados_raw = carregar_dados(file_reativados)
            df_fortpel_raw = carregar_dados(file_fortpel)
            df_adm_raw = carregar_dados(file_adm)
            df_categorias_raw = carregar_dados(file_categorias) if file_categorias else None
    except Exception as e:
        st.error(f"Erro ao ler arquivos: {e}")
        st.stop()

    qtd_bases = 5 if df_categorias_raw is not None else 4
    st.success(f"As {qtd_bases} bases foram carregadas com sucesso!")

    st.subheader("2. Mapeamento das Colunas Chave")
    titulos_abas_map = ["Bases do BI", "Acompanhamento Fortpel", "Planilha ADM (Orçamentos)"]
    if df_categorias_raw is not None:
        titulos_abas_map.append("Planilha de Categorias (BI)")

    abas_mapeamento = st.tabs(titulos_abas_map)
    tab_bi = abas_mapeamento[0]
    tab_fortpel = abas_mapeamento[1]
    tab_adm = abas_mapeamento[2]
    tab_cat = abas_mapeamento[3] if df_categorias_raw is not None else None

    with tab_bi:
        c_b1, c_b2 = st.columns(2)
        with c_b1:
            col_novos_cli = st.selectbox(
                "Coluna Cliente (BI Novos)",
                df_novos_raw.columns,
                index=df_novos_raw.columns.get_loc(identificar_coluna(df_novos_raw.columns, ["cliente", "razao", "nome"]))
            )
            col_novos_val = identificar_coluna(df_novos_raw.columns, ["venda", "valor", "total", "faturado"])
        with c_b2:
            col_reat_cli = st.selectbox(
                "Coluna Cliente (BI Reativados)",
                df_reativados_raw.columns,
                index=df_reativados_raw.columns.get_loc(identificar_coluna(df_reativados_raw.columns, ["cliente", "razao", "nome"]))
            )
            col_reat_val = identificar_coluna(df_reativados_raw.columns, ["venda", "valor", "total", "faturado"])

    with tab_fortpel:
        c_f1, c_f2, c_f3 = st.columns(3)
        c_f4, c_f5, c_f6 = st.columns(3)
        with c_f1:
            col_f_cli = st.selectbox("Cliente", df_fortpel_raw.columns, index=df_fortpel_raw.columns.get_loc(identificar_coluna(df_fortpel_raw.columns, ["cliente", "razao", "nome"])))
        with c_f2:
            col_f_ped = st.selectbox("Pedido", df_fortpel_raw.columns, index=df_fortpel_raw.columns.get_loc(identificar_coluna(df_fortpel_raw.columns, ["pedido", "numero", "ordem"])))
        with c_f3:
            idx_not = df_fortpel_raw.columns.get_loc(identificar_coluna(df_fortpel_raw.columns, ["nota", "nf", "n_nota", "n_nota_obs", "pedido"]))
            col_f_nota = st.selectbox("Nota Fiscal (para Maior Venda em Nota Única):", df_fortpel_raw.columns, index=idx_not)
        with c_f4:
            col_f_vend = st.selectbox("Vendedor", df_fortpel_raw.columns, index=df_fortpel_raw.columns.get_loc(identificar_coluna(df_fortpel_raw.columns, ["carteira", "operador", "vendedor", "representante", "gestor"])))
        with c_f5:
            col_f_val = st.selectbox("Valor", df_fortpel_raw.columns, index=df_fortpel_raw.columns.get_loc(identificar_coluna(df_fortpel_raw.columns, ["valor", "total", "faturado", "preco", "somadevl"])))
        with c_f6:
            col_f_data = st.selectbox("Data", df_fortpel_raw.columns, index=df_fortpel_raw.columns.get_loc(identificar_coluna(df_fortpel_raw.columns, ["dt_inser", "dt inser", "data inser", "dt_saida", "data", "emissao", "faturamento", "dt"])))

        # Opção inteligente para valor da Aba 1 (caso Fortpel tenha apenas soma total)
        tem_valor_bi = (col_novos_val is not None) and (col_reat_val is not None)
        origem_valor_aba1 = st.radio(
            "Origem do Valor da Aba 1 (Novos/Reativados):",
            [
                "Usar coluna 'Venda' do BI (Novos e Reativados) [Recomendado se o Acompanhamento tiver total geral]",
                "Usar coluna selecionada do Acompanhamento Fortpel"
            ] if tem_valor_bi else ["Usar coluna selecionada do Acompanhamento Fortpel"],
            index=0 if tem_valor_bi else 0
        )

    with tab_adm:
        st.markdown("📊 **Tabela Dinâmica ADM (Aba 2: Propostas Enviadas)**")
        st.caption("A aba de propostas é gerada pela contagem de registros do campo Status por Vendedor.")
        c_a1, c_a2 = st.columns(2)
        with c_a1:
            idx_v = df_adm_raw.columns.get_loc(identificar_coluna(df_adm_raw.columns, ["vendedor", "operador", "usuario", "representante"]))
            col_a_vend = st.selectbox("Coluna Vendedor (ADM):", df_adm_raw.columns, index=idx_v)
        with c_a2:
            idx_s = df_adm_raw.columns.get_loc(identificar_coluna(df_adm_raw.columns, ["status", "situacao", "fase"]))
            col_a_status = st.selectbox("Coluna Status / Contagem (ADM):", df_adm_raw.columns, index=idx_s)

        todos_status_adm = sorted([str(s) for s in df_adm_raw[col_a_status].dropna().unique()])
        status_adm_selecionados = st.multiselect(
            "Filtrar por Status específicos (opcional - por padrão conta todos):",
            options=todos_status_adm,
            default=todos_status_adm,
            help="Selecione quais status entram na contagem. Por padrão, todos os status presentes na planilha são contabilizados."
        )

    if tab_cat is not None:
        with tab_cat:
            st.markdown("📦 **Planilha de Categorias do BI**")
            st.caption("Verificação automática das categorias oficiais na planilha carregada.")
            encontradas = [cat for cat in CATEGORIAS_ROLETA if any(normalizar_texto(cat).replace("'", "").replace(" ", "") in normalizar_texto(str(col)).replace("'", "").replace(" ", "") for col in df_categorias_raw.columns)]
            if encontradas:
                st.success(f"✅ {len(encontradas)} das 7 categorias oficiais identificadas nas colunas: {', '.join(encontradas)}")
            else:
                st.info("Formato detalhado por produto/linha detectado. O sistema cruzará os dados automaticamente.")

    # Conversão de Datas (Apenas Fortpel requer data)
    df_fortpel = df_fortpel_raw.copy()
    df_fortpel['__DATA_DT'] = pd.to_datetime(df_fortpel[col_f_data], dayfirst=True, errors='coerce')

    datas_validas = df_fortpel['__DATA_DT'].dropna()
    min_date = datas_validas.min() if not datas_validas.empty else pd.Timestamp.now()
    max_date = datas_validas.max() if not datas_validas.empty else pd.Timestamp.now()

    # --- FILTRO POR PERÍODO ---
    st.subheader("3. Período e Filtro de Datas")

    c_filt1, c_filt2 = st.columns([1, 1])
    with c_filt1:
        modo_filtro_roleta = st.radio(
            "Filtro de Data para a Roleta (Aba 1 - Novos e Reativados):",
            [
                "Opção B: Filtrar pelo intervalo de datas selecionado (Inclusivo, recomendado)",
                "Opção A: Considerar tudo que vier no BI (Sem restrição de data na Roleta)"
            ],
            index=0,
            help="A Opção B restringe os pedidos rigorosamente entre as datas selecionadas, evitando trazer pedidos inseridos após a data limite ou pedidos antigos."
        )

        criterio_roleta = st.radio(
            "Critério de Contagem da Roleta (Aba 1):",
            [
                "1 Registro por Cliente Único (Padrão BI - 1 linha por cliente, soma exata)",
                "Todos os Pedidos Individuais (Pode gerar múltiplas linhas se o cliente fez mais de um pedido)"
            ],
            index=0,
            help="No relatório do BI, cada cliente conta como 1 registro. Marque '1 Registro por Cliente Único' para bater 100% com o relatório do BI."
        )

    with c_filt2:
        data_sel = st.date_input(
            "Selecione o intervalo de datas (para Aba 2 e filtro da Roleta):",
            value=(min_date.date(), max_date.date())
        )

    if isinstance(data_sel, tuple) and len(data_sel) == 2:
        dt_inicio, dt_fim = data_sel

        if st.button("🚀 Processar e Gerar Planilha Completa", type="primary"):
            with st.spinner("Processando cruzamentos das bases..."):
                # ==========================================
                # ABA 1: REATIVADOS - NOVOS
                # ==========================================
                df_novos_clean = df_novos_raw.dropna(subset=[col_novos_cli]).copy()
                df_novos_clean['__CLI_NORM'] = df_novos_clean[col_novos_cli].apply(normalizar_texto)
                clientes_novos = set(df_novos_clean['__CLI_NORM'])

                df_reat_clean = df_reativados_raw.dropna(subset=[col_reat_cli]).copy()
                df_reat_clean['__CLI_NORM'] = df_reat_clean[col_reat_cli].apply(normalizar_texto)
                clientes_reativados = set(df_reat_clean['__CLI_NORM'])

                df_fortpel['__CLI_NORM'] = df_fortpel[col_f_cli].apply(normalizar_texto)

                def atribuir_status(nome_norm):
                    if nome_norm in clientes_novos:
                        return "NOVOS"
                    elif nome_norm in clientes_reativados:
                        return "REATIVADO"
                    return None

                df_fortpel['STATUS'] = df_fortpel['__CLI_NORM'].apply(atribuir_status)

                status_validos = ['NOVOS', 'NOVO', 'REATIVADO', 'REATIVADOS']
                mask_status = df_fortpel['STATUS'].astype(str).str.upper().isin(status_validos)

                if "Sem restrição" in modo_filtro_roleta:
                    mask_data_f = (df_fortpel['__DATA_DT'].dt.date <= dt_fim) if pd.notna(dt_fim) else True
                    df_aba1 = df_fortpel[mask_data_f & mask_status].copy()
                else:
                    mask_data_f = (df_fortpel['__DATA_DT'].dt.date >= dt_inicio) & (df_fortpel['__DATA_DT'].dt.date <= dt_fim)
                    df_aba1 = df_fortpel[mask_data_f & mask_status].copy()

                    if df_aba1.empty and mask_status.sum() > 0:
                        mask_data_f_flex = (df_fortpel['__DATA_DT'].dt.date <= dt_fim) if pd.notna(dt_fim) else True
                        df_aba1 = df_fortpel[mask_data_f_flex & mask_status].copy()
                        if not df_aba1.empty:
                            st.warning(f"⚠️ Atenção: Não foram encontrados pedidos exatamente no intervalo selecionado ({dt_inicio.strftime('%d/%m/%Y')} a {dt_fim.strftime('%d/%m/%Y')}), mas os clientes foram localizados na base. Os pedidos foram recuperados considerando os registros até a data limite!")

                if "Cliente Único" in criterio_roleta:
                    df_aba1 = df_aba1.sort_values(by=['__DATA_DT', col_f_ped], ascending=[True, True])
                    df_aba1 = df_aba1.drop_duplicates(subset=['__CLI_NORM'], keep='first')

                if "BI" in origem_valor_aba1:
                    mapa_valores_bi = {}
                    if col_novos_val:
                        for _, row in df_novos_clean.iterrows():
                            mapa_valores_bi[row['__CLI_NORM']] = limpar_valor(row[col_novos_val])
                    if col_reat_val:
                        for _, row in df_reat_clean.iterrows():
                            mapa_valores_bi[row['__CLI_NORM']] = limpar_valor(row[col_reat_val])
                    df_aba1['VALOR_NUM'] = df_aba1['__CLI_NORM'].map(mapa_valores_bi).fillna(0.0)
                else:
                    df_aba1['VALOR_NUM'] = df_aba1[col_f_val].apply(limpar_valor)

                # Regra Duas Compras
                meses_periodo = set(pd.date_range(dt_inicio, dt_fim).to_period('M'))
                mask_meses_f = df_fortpel['__DATA_DT'].dt.to_period('M').isin(meses_periodo)
                pedidos_no_mes = df_fortpel[mask_meses_f].groupby('__CLI_NORM')[col_f_ped].nunique().to_dict()

                def verificar_duas_compras(cli_norm):
                    qtd = pedidos_no_mes.get(cli_norm, 1)
                    return qtd > 1

                def gerar_rotulo_brinde(cli_norm):
                    qtd = pedidos_no_mes.get(cli_norm, 1)
                    if qtd == 2:
                        return "DUAS COMPRAS"
                    elif qtd > 2:
                        return f"DUAS COMPRAS ({qtd} PEDIDOS)"
                    return ""

                df_aba1['TEM_DUAS_COMPRAS'] = df_aba1['__CLI_NORM'].apply(verificar_duas_compras)
                df_aba1['BRINDE ROLETA'] = df_aba1['__CLI_NORM'].apply(gerar_rotulo_brinde)

                # ==========================================
                # ABA 2: PROPOSTAS ENVIADAS (TABELA DINÂMICA ADM)
                # ==========================================
                df_adm_filtrado = df_adm_raw.dropna(subset=[col_a_vend]).copy()
                if status_adm_selecionados:
                    df_adm_filtrado = df_adm_filtrado[df_adm_filtrado[col_a_status].astype(str).isin(status_adm_selecionados)]

                df_aba2 = (
                    df_adm_filtrado.groupby(col_a_vend)[col_a_status]
                    .count()
                    .reset_index()
                )
                df_aba2.columns = ['Vendedor', 'Qtd']
                df_aba2 = df_aba2.sort_values(by=['Qtd', 'Vendedor'], ascending=[False, True])

                # ========================================================
                # 🏆 MAIOR VENDA EM NOTA ÚNICA (EXTRAÍDA DAS OUTRAS BASES)
                # ========================================================
                df_top_notas = extrair_maior_venda_nota(
                    df_fortpel=df_fortpel,
                    col_nota=col_f_nota,
                    col_vend=col_f_vend,
                    col_val=col_f_val,
                    dt_inicio=dt_inicio,
                    dt_fim=dt_fim,
                    df_cat_raw=df_categorias_raw
                )

                # ========================================================
                # 🏷️ CATEGORIAS DO BI (SE PLANILHA 5 FOI CARREGADA)
                # ========================================================
                df_rank_cat = pd.DataFrame()
                total_geral_cat = 0.0
                df_lideres_cat = pd.DataFrame()

                if df_categorias_raw is not None:
                    df_rank_cat, total_geral_cat, df_lideres_cat = extrair_ranking_categoria(
                        df_cat_raw=df_categorias_raw,
                        categoria_alvo=categoria_solicitada,
                        df_fortpel=df_fortpel,
                        df_adm=df_adm_raw
                    )

                # ==========================================
                # GERAÇÃO DO WORKBOOK EXCEL (3 ABAS)
                # ==========================================
                wb = openpyxl.Workbook()

                dia_i = dt_inicio.strftime('%d')
                dia_f = dt_fim.strftime('%d')
                nome_aba1 = f"{dia_i}-{dia_f} Reativados - Novos"[:31]
                nome_aba2 = f"{dia_i}-{dia_f} Propostas Enviadas"[:31]

                # Preencher Aba 1
                ws1 = wb.active
                ws1.title = nome_aba1
                if not df_aba1.empty:
                    df_aba1['VENDEDOR_CLEAN'] = df_aba1[col_f_vend].astype(str).str.strip().str.upper()
                    df_aba1 = df_aba1.sort_values(by=['VENDEDOR_CLEAN', '__DATA_DT', col_f_cli])
                    estilizar_aba(ws1, df_aba1, col_f_ped, col_f_val, col_f_cli, '__DATA_DT', titulo_col_id="PEDIDO")
                else:
                    ws1.cell(row=1, column=1, value="Nenhum pedido encontrado para Novos/Reativados no período.")

                # Preencher Aba 2
                ws2 = wb.create_sheet(title=nome_aba2)
                if not df_aba2.empty:
                    estilizar_aba_propostas(ws2, df_aba2)
                else:
                    ws2.cell(row=1, column=1, value="Nenhum registro encontrado na planilha ADM.")

                # Preencher Aba 3: Destaques Roleta (Nota Única + Categorias)
                ws3 = wb.create_sheet(title="Destaques Roleta")
                estilizar_aba_produtos(ws3, df_top_notas, df_lideres_cat)

                nome_arquivo_saida = f"Controle_roleta_{dt_inicio.strftime('%d%m')}_a_{dt_fim.strftime('%d%m')}.xlsx"
                caminho_salvar_local = PASTA_SAIDA / nome_arquivo_saida
                wb.save(caminho_salvar_local)

                buffer_download = io.BytesIO()
                wb.save(buffer_download)
                buffer_download.seek(0)

                # Salva no Session State para manter persistente entre seleções
                st.session_state['processado'] = True
                st.session_state['df_aba1'] = df_aba1
                st.session_state['df_aba2'] = df_aba2
                st.session_state['df_top_notas'] = df_top_notas
                st.session_state['df_rank_cat'] = df_rank_cat
                st.session_state['total_geral_cat'] = total_geral_cat
                st.session_state['df_lideres_cat'] = df_lideres_cat
                st.session_state['df_categorias_raw'] = df_categorias_raw
                st.session_state['df_fortpel'] = df_fortpel
                st.session_state['df_adm_raw'] = df_adm_raw
                st.session_state['categoria_solicitada'] = categoria_solicitada
                st.session_state['col_f_ped'] = col_f_ped
                st.session_state['col_f_cli'] = col_f_cli
                st.session_state['nome_aba1'] = nome_aba1
                st.session_state['nome_aba2'] = nome_aba2
                st.session_state['caminho_salvar_local'] = caminho_salvar_local
                st.session_state['buffer_bytes'] = buffer_download.getvalue()
                st.session_state['nome_arquivo_saida'] = nome_arquivo_saida
                st.session_state['tem_categorias'] = (df_categorias_raw is not None)

        # Exibição dos Resultados (Persistente no Session State)
        if st.session_state.get('processado', False):
            df_aba1_s = st.session_state['df_aba1']
            df_aba2_s = st.session_state['df_aba2']
            nome_aba1_s = st.session_state['nome_aba1']
            nome_aba2_s = st.session_state['nome_aba2']
            caminho_local_s = st.session_state['caminho_salvar_local']
            buffer_bytes_s = st.session_state['buffer_bytes']
            nome_saida_s = st.session_state['nome_arquivo_saida']

            st.success("✅ Planilha gerada com 3 abas consolidadas com sucesso!")
            st.info(f"📁 Arquivo salvo em: `{caminho_local_s}`")

            abas_titulos = [f"Aba 1: {nome_aba1_s}", f"Aba 2: {nome_aba2_s}", "Aba 3: 🏆 Destaques Comerciais (Roleta)"]
            abas_res = st.tabs(abas_titulos)

            with abas_res[0]:
                if not df_aba1_s.empty:
                    qtd_duas_compras = int(df_aba1_s['TEM_DUAS_COMPRAS'].sum()) if 'TEM_DUAS_COMPRAS' in df_aba1_s else 0
                    if qtd_duas_compras > 0:
                        st.warning(f"⚠️ {qtd_duas_compras} cliente(s) com mais de uma compra no mesmo mês identificado(s) (destacados com 'DUAS COMPRAS' e linha em vermelho na planilha)!")
                    st.dataframe(df_aba1_s[['VENDEDOR_CLEAN', st.session_state['col_f_ped'], st.session_state['col_f_cli'], 'VALOR_NUM', 'STATUS', '__DATA_DT', 'BRINDE ROLETA']], use_container_width=True)
                else:
                    st.write("Sem registros no período para a Aba 1.")

            with abas_res[1]:
                if not df_aba2_s.empty:
                    st.dataframe(df_aba2_s, use_container_width=True)
                    st.metric("Total de Propostas Efetivadas", int(df_aba2_s['Qtd'].sum()))
                else:
                    st.write("Sem registros na planilha ADM.")

            with abas_res[2]:
                df_top_notas_s = st.session_state.get('df_top_notas', pd.DataFrame())
                df_cat_raw_s = st.session_state.get('df_categorias_raw')
                df_fort_s = st.session_state.get('df_fortpel')
                df_adm_s = st.session_state.get('df_adm_raw')
                cat_padrao = st.session_state.get('categoria_solicitada', 'HIGIENE E LIMPEZA')

                # ========================================================
                # 🏆 1. MAIOR VENDA EM NOTA ÚNICA (INFORMAÇÃO DA FOTO)
                # ========================================================
                st.markdown("### 🏆 Maior Venda em Nota Única")
                st.caption("Extraído das bases de acompanhamento e pedidos carregadas.")

                if not df_top_notas_s.empty:
                    top1_nota = df_top_notas_s.iloc[0]
                    nota_formatada = int(top1_nota['Nota']) if isinstance(top1_nota['Nota'], (int, float)) and top1_nota['Nota'] == int(top1_nota['Nota']) else str(top1_nota['Nota'])
                    val_formatado = f"R$ {top1_nota['Valor_Total']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

                    c_m1, c_m2, c_m3 = st.columns(3)
                    with c_m1:
                        st.metric("Vendedor Campeão (Nota Única)", str(top1_nota['Vendedor']))
                    with c_m2:
                        st.metric("Nota Fiscal", nota_formatada)
                    with c_m3:
                        st.metric("Valor Total da Nota", val_formatado)

                    with st.expander("📋 Ver Ranking das Maiores Vendas em Nota Única"):
                        df_notas_disp = df_top_notas_s.copy()
                        df_notas_disp['Valor_Total'] = df_notas_disp['Valor_Total'].apply(lambda x: f"R$ {x:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
                        st.dataframe(df_notas_disp, use_container_width=True)
                else:
                    st.info("Informações de nota única disponíveis na base de Acompanhamento Fortpel.")

                st.markdown("---")

                # ========================================================
                # 🏷️ 2. MAIOR VENDA POR CATEGORIA SOLICITADA
                # ========================================================
                st.markdown("### 🏷️ Consulta de Maior Venda por Categoria Solicitada")
                st.caption("Selecione a categoria desejada para visualizar o vendedor que fez a maior venda:")

                idx_cat_sel = CATEGORIAS_ROLETA.index(cat_padrao) if cat_padrao in CATEGORIAS_ROLETA else 0
                cat_escolhida = st.selectbox(
                    "Selecione a Categoria:",
                    CATEGORIAS_ROLETA,
                    index=idx_cat_sel,
                    key="sel_cat_destaque"
                )

                if df_cat_raw_s is not None:
                    # Calcula instantaneamente para a categoria escolhida no dropdown
                    df_rank_escolhido, total_escolhido, df_lideres_todas = extrair_ranking_categoria(
                        df_cat_raw=df_cat_raw_s,
                        categoria_alvo=cat_escolhida,
                        df_fortpel=df_fort_s,
                        df_adm=df_adm_s
                    )

                    if not df_rank_escolhido.empty:
                        lider_cat = df_rank_escolhido.iloc[0]
                        val_lider_fmt = f"R$ {lider_cat['Total_Vendido']:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
                        total_geral_fmt = f"R$ {total_escolhido:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

                        col_c1, col_c2, col_c3 = st.columns(3)
                        with col_c1:
                            st.metric("Vendedor Líder da Categoria", str(lider_cat['Vendedor']))
                        with col_c2:
                            st.metric(f"Vendido por {lider_cat['Vendedor']}", val_lider_fmt)
                        with col_c3:
                            st.metric(f"Total Geral ({cat_escolhida})", total_geral_fmt)

                        st.markdown(f"**Ranking de Vendedores na categoria `{cat_escolhida}`:**")
                        df_cat_disp = df_rank_escolhido.copy()
                        df_cat_disp['Total_Vendido'] = df_cat_disp['Total_Vendido'].apply(lambda x: f"R$ {x:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
                        st.dataframe(df_cat_disp, use_container_width=True)

                    if not df_lideres_todas.empty:
                        with st.expander("📊 Ver Tabela Geral de Líderes das 7 Categorias"):
                            df_lid_disp = df_lideres_todas.copy()
                            df_lid_disp['Valor_Total'] = df_lid_disp['Valor_Total'].apply(lambda x: f"R$ {x:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
                            df_lid_disp['Total_Categoria'] = df_lid_disp['Total_Categoria'].apply(lambda x: f"R$ {x:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'))
                            st.dataframe(df_lid_disp, use_container_width=True)
                else:
                    st.info("ℹ️ Para consultar os vendedores líderes por categoria, faça o upload da Planilha de Categorias do BI no Campo 5.")

            st.download_button(
                label="📥 Baixar Planilha Consolidada (3 Abas) (.xlsx)",
                data=buffer_bytes_s,
                file_name=nome_saida_s,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
